"""Import and edit the tenant registry from CSV and Excel workbooks."""

import csv
import os
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from rental_automation.models import Tenant, TenantIdentifier

TENANT_WORKBOOK_PATH = Path(__file__).resolve().parent.parent / "tenants.xlsx"
TENANT_WORKSHEET = "Tenants"
TENANT_COLUMNS = (
    "name",
    "unit_number",
    "phone",
    "email",
    "upi_vpa",
    "monthly_rent",
    "maintenance_fee",
    "deposit_amount",
    "opening_balance",
    "due_day",
    "is_active",
    "expected_sender_name",
    "expected_remarks",
)


def import_tenants(session: Session, source: Path, *, dry_run: bool = False) -> int:
    """Upsert tenants by unit number from a UTF-8 CSV file."""
    with source.open("r", encoding="utf-8-sig", newline="") as file:
        rows = (
            (_normalize_row(row), line_number)
            for line_number, row in enumerate(csv.DictReader(file), start=2)
        )
        return _import_rows(session, rows, "CSV", dry_run=dry_run)


def import_tenants_workbook(
    session: Session,
    source: Path = TENANT_WORKBOOK_PATH,
    *,
    dry_run: bool = False,
    commit: bool = True,
) -> int:
    """Upsert tenant records from the workbook's ``Tenants`` worksheet."""
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not open tenant workbook {source}: {exc}") from exc

    try:
        if TENANT_WORKSHEET not in workbook.sheetnames:
            raise ValueError(f"Tenant workbook must contain a '{TENANT_WORKSHEET}' sheet")
        sheet = workbook[TENANT_WORKSHEET]
        heading_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headings = [_normalize_heading(value) for value in heading_row]
        missing = {"name", "unit_number", "phone", "monthly_rent"} - set(headings)
        if missing:
            raise ValueError(
                "Tenant workbook is missing required columns: "
                + ", ".join(sorted(missing))
            )

        try:
            workbook_rows = []
            seen_units = set()
            for line_number, cells in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                raw = dict(zip(headings, cells))
                if not any(value not in (None, "") for value in raw.values()):
                    continue
                unit_number = _cell_text(raw.get("unit_number"))
                if unit_number and unit_number in seen_units:
                    raise ValueError(
                        f"Duplicate unit_number '{unit_number}' on row {line_number}"
                    )
                if unit_number:
                    seen_units.add(unit_number)
                workbook_rows.append((
                    {
                        key: _cell_text(raw.get(key))
                        for key in set(TENANT_COLUMNS) | set(headings)
                    },
                    line_number,
                ))
        except Exception as exc:
            raise ValueError(f"Could not read tenant workbook {source}: {exc}") from exc

        return _import_rows(
            session,
            workbook_rows,
            "workbook",
            dry_run=dry_run,
            commit=commit,
            replace_identifiers=True,
        )
    finally:
        workbook.close()


def upsert_tenant_workbook_row(values: dict, source: Path = TENANT_WORKBOOK_PATH) -> None:
    """Add or update one tenant in the workbook, saving it atomically."""
    try:
        workbook = load_workbook(source)
    except Exception as exc:
        raise ValueError(f"Could not open tenant workbook {source}: {exc}") from exc
    try:
        if TENANT_WORKSHEET not in workbook.sheetnames:
            raise ValueError(f"Tenant workbook must contain a '{TENANT_WORKSHEET}' sheet")
        sheet = workbook[TENANT_WORKSHEET]
        headings = {
            _normalize_heading(cell.value): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        missing = set(TENANT_COLUMNS) - set(headings)
        if missing:
            raise ValueError(
                "Tenant workbook is missing columns: " + ", ".join(sorted(missing))
            )

        unit_number = str(values["unit_number"]).strip()
        target_row = None
        first_empty_row = None
        for row_number in range(2, sheet.max_row + 1):
            existing_unit = sheet.cell(row_number, headings["unit_number"]).value
            if existing_unit is not None and str(existing_unit).strip() == unit_number:
                target_row = row_number
                break
            if first_empty_row is None and all(
                sheet.cell(row_number, column).value in (None, "")
                for column in headings.values()
            ):
                first_empty_row = row_number
        if target_row is None:
            target_row = first_empty_row or max(sheet.max_row + 1, 2)

        record = {key: values.get(key) for key in TENANT_COLUMNS}
        record["is_active"] = values.get("is_active", True)
        for key, column in headings.items():
            if key in record:
                sheet.cell(target_row, column).value = record[key]
        _format_tenant_row(sheet, target_row, headings)
        _resize_tenant_table(sheet, target_row, max(headings.values()))
        _save_workbook(workbook, source)
    finally:
        workbook.close()


def export_tenants_workbook(session: Session, source: Path = TENANT_WORKBOOK_PATH) -> int:
    """Write the current database tenant registry back to the workbook."""
    source = Path(source)
    source.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook = load_workbook(source) if source.exists() else _new_registry_workbook()
    except Exception as exc:
        raise ValueError(f"Could not open tenant workbook {source}: {exc}") from exc
    try:
        if TENANT_WORKSHEET not in workbook.sheetnames:
            raise ValueError(f"Tenant workbook must contain a '{TENANT_WORKSHEET}' sheet")
        sheet = workbook[TENANT_WORKSHEET]
        headings = {
            _normalize_heading(cell.value): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        missing = set(TENANT_COLUMNS) - set(headings)
        if missing:
            raise ValueError(
                "Tenant workbook is missing columns: " + ", ".join(sorted(missing))
            )

        tenants = list(session.scalars(select(Tenant).order_by(Tenant.unit_number)))
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for tenant in tenants:
            record = {
                "name": tenant.name,
                "unit_number": tenant.unit_number,
                "phone": tenant.phone,
                "email": tenant.email,
                "upi_vpa": tenant.upi_vpa,
                "monthly_rent": tenant.monthly_rent,
                "maintenance_fee": tenant.maintenance_fee,
                "deposit_amount": tenant.deposit_amount,
                "opening_balance": tenant.opening_balance,
                "due_day": tenant.due_day,
                "is_active": tenant.is_active,
                "expected_sender_name": next(
                    (item.value for item in tenant.identifiers if item.kind == "sender_name"),
                    None,
                ),
                "expected_remarks": next(
                    (item.value for item in tenant.identifiers if item.kind == "remarks"),
                    None,
                ),
            }
            row_number = sheet.max_row + 1
            for key, column in headings.items():
                if key in record:
                    sheet.cell(row_number, column).value = record[key]
            _format_tenant_row(sheet, row_number, headings)
        _resize_tenant_table(sheet, sheet.max_row, max(headings.values()))
        _save_workbook(workbook, source)
        return len(tenants)
    finally:
        workbook.close()


def _new_registry_workbook() -> Workbook:
    """Create an empty workbook when a deployment has no local data file."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TENANT_WORKSHEET
    sheet.append(TENANT_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Tenant workbook"])
    instructions.append(["Add or update tenants on the Tenants sheet using unit_number."])
    instructions.append(["For cloud deployments, upload the edited workbook in the admin dashboard."])
    return workbook


def _import_rows(
    session: Session,
    rows,
    source_label: str,
    *,
    dry_run: bool,
    commit: bool = True,
    replace_identifiers: bool = False,
) -> int:
    count = 0
    for values, line_number in rows:
        try:
            unit_number = _required(values, "unit_number")
            name = _required(values, "name")
            phone = _required(values, "phone")
            rent = _money(_required(values, "monthly_rent"), "monthly_rent")
            maintenance = _money(values.get("maintenance_fee") or "0", "maintenance_fee")
            deposit = _money(values.get("deposit_amount") or "0", "deposit_amount")
            opening = _money(values.get("opening_balance") or "0", "opening_balance")
            due_day = _integer(values.get("due_day"), "due_day")
            if due_day is not None and not 1 <= due_day <= 31:
                raise ValueError("due_day must be between 1 and 31")
            is_active = _boolean(values.get("is_active", "true"))
        except (ValueError, InvalidOperation) as exc:
            raise ValueError(f"Invalid tenant {source_label} row {line_number}: {exc}") from exc

        tenant = session.scalar(select(Tenant).where(Tenant.unit_number == unit_number))
        if tenant is None:
            tenant = Tenant(unit_number=unit_number)
            session.add(tenant)

        tenant.name = name
        tenant.phone = phone
        tenant.monthly_rent = rent
        tenant.maintenance_fee = maintenance
        tenant.deposit_amount = deposit
        tenant.opening_balance = opening
        tenant.due_day = due_day
        tenant.email = values.get("email") or None
        tenant.upi_vpa = values.get("upi_vpa") or None
        tenant.is_active = is_active

        for kind, column in (
            ("sender_name", "expected_sender_name"),
            ("remarks", "expected_remarks"),
        ):
            expected = values.get(column)
            if replace_identifiers:
                for identifier in list(tenant.identifiers):
                    if identifier.kind == kind and identifier.value != expected:
                        tenant.identifiers.remove(identifier)
            if expected and not any(
                identifier.kind == kind and identifier.value == expected
                for identifier in tenant.identifiers
            ):
                tenant.identifiers.append(TenantIdentifier(kind=kind, value=expected))
        count += 1

    if not dry_run and commit:
        session.commit()
    elif dry_run:
        session.rollback()
    return count


def _required(values: dict[str, str], key: str) -> str:
    value = values.get(key, "")
    if not value:
        raise ValueError(f"missing required column value {key}")
    return value


def _money(value: str, field: str) -> Decimal:
    amount = Decimal(value.replace(",", "").replace(chr(0x20B9), "").strip())
    if field != "opening_balance" and amount < 0:
        raise ValueError(f"{field} cannot be negative")
    return amount


def _boolean(value: str) -> bool:
    normalized = value.casefold()
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise ValueError("is_active must be true/false, yes/no, or 1/0")


def _integer(value: str, field: str) -> int | None:
    if not value:
        return None
    number = Decimal(value)
    if number != number.to_integral_value():
        raise ValueError(f"{field} must be a whole number")
    return int(number)


def _normalize_row(row: dict) -> dict[str, str]:
    return {
        _normalize_heading(key): _cell_text(value)
        for key, value in row.items()
        if key
    }


def _normalize_heading(value) -> str:
    return str(value or "").strip().casefold().replace(" ", "_")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _resize_tenant_table(sheet, last_row: int, last_column: int) -> None:
    for table in sheet.tables.values():
        table.ref = f"A1:{_excel_column(last_column)}{max(2, last_row)}"


def _format_tenant_row(sheet, row_number: int, headings: dict[str, int]) -> None:
    for key in ("unit_number", "phone", "email", "upi_vpa"):
        if key in headings:
            sheet.cell(row_number, headings[key]).number_format = "@"
    for key in ("monthly_rent", "maintenance_fee", "deposit_amount", "opening_balance"):
        if key in headings:
            sheet.cell(row_number, headings[key]).number_format = "#,##0.00"
    if "due_day" in headings:
        sheet.cell(row_number, headings["due_day"]).number_format = "0"


def _save_workbook(workbook, source: Path) -> None:
    temporary_path: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix="tenants-", suffix=".xlsx", dir=source.parent, delete=False
        ) as temporary_file:
            temporary_path = temporary_file.name
        workbook.save(temporary_path)
        os.replace(temporary_path, source)
        temporary_path = None
    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)


def _excel_column(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result
