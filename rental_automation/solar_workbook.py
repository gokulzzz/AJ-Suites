"""Excel reporting for per-consumer solar readings and KSEB bills."""

from __future__ import annotations

import os
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session

from rental_automation.models import KSEBConsumerMonthlyBill

SOLAR_WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "solar_consumption.xlsx"

_HEADERS = [
    "Billing Month",
    "Consumer Number",
    "Bill Date",
    "Imported Normal (kWh)",
    "Imported Off Peak (kWh)",
    "Imported Peak (kWh)",
    "Imported Total (kWh)",
    "Exported Normal (kWh)",
    "Exported Off Peak (kWh)",
    "Exported Peak (kWh)",
    "Exported Total (kWh)",
    "Solar Normal (kWh)",
    "Solar Off Peak (kWh)",
    "Solar Peak (kWh)",
    "Solar Generation (kWh)",
    "Solar Self-Consumed (kWh)",
    "Site Units Consumed (kWh)",
    "KSEB Bill Payable (INR)",
    "Data Source",
    "Last Fetched",
]


def _style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244536")
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _as_excel_number(value) -> float:
    return float(value or 0)


def export_solar_workbook(
    session: Session, destination: Path = SOLAR_WORKBOOK_PATH
) -> Path:
    """Rebuild the workbook from saved service bills and atomically replace it."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    bills = list(
        session.scalars(
            select(KSEBConsumerMonthlyBill).order_by(
                KSEBConsumerMonthlyBill.billing_month,
                KSEBConsumerMonthlyBill.consumer_number,
            )
        )
    )

    workbook = Workbook()
    usage = workbook.active
    usage.title = "Solar Usage"
    usage.append(_HEADERS)
    months: dict[object, set[str]] = defaultdict(set)

    for bill in bills:
        row_number = usage.max_row + 1
        usage.append(
            [
                bill.billing_month,
                bill.consumer_number,
                bill.billing_date,
                _as_excel_number(bill.imported_normal_kwh),
                _as_excel_number(bill.imported_off_peak_kwh),
                _as_excel_number(bill.imported_peak_kwh),
                f"=SUM(D{row_number}:F{row_number})",
                _as_excel_number(bill.exported_normal_kwh),
                _as_excel_number(bill.exported_off_peak_kwh),
                _as_excel_number(bill.exported_peak_kwh),
                f"=SUM(H{row_number}:J{row_number})",
                _as_excel_number(bill.solar_normal_kwh),
                _as_excel_number(bill.solar_off_peak_kwh),
                _as_excel_number(bill.solar_peak_kwh),
                f"=SUM(L{row_number}:N{row_number})",
                f"=O{row_number}-K{row_number}",
                f"=O{row_number}+G{row_number}-K{row_number}",
                _as_excel_number(bill.bill_amount),
                bill.source_label,
                bill.fetched_at.replace(tzinfo=None),
            ]
        )
        months[bill.billing_month].add(bill.consumer_number)
        usage.cell(row_number, 1).number_format = "mmm yyyy"
        usage.cell(row_number, 3).number_format = "dd mmm yyyy"
        usage.cell(row_number, 20).number_format = "dd mmm yyyy hh:mm"
        for column in range(4, 18):
            usage.cell(row_number, column).number_format = "#,##0.000"
        usage.cell(row_number, 18).number_format = '₹ #,##0.00'

    _style_sheet(
        usage,
        {
            "A": 16,
            "B": 19,
            "C": 16,
            **{chr(64 + index): 20 for index in range(4, 19)},
            "S": 17,
            "T": 22,
        },
    )
    if bills:
        table = Table(displayName="SolarUsage", ref=f"A1:T{usage.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        usage.add_table(table)

    summary = workbook.create_sheet("Monthly Summary")
    summary.append(
        [
            "Billing Month",
            "Connected Services",
            "Grid Imported (kWh)",
            "Grid Exported (kWh)",
            "Solar Generation (kWh)",
            "Solar Self-Consumed (kWh)",
            "Site Units Consumed (kWh)",
            "KSEB Bill Payable (INR)",
        ]
    )
    for row_number, billing_month in enumerate(sorted(months), start=2):
        summary.append(
            [
                billing_month,
                len(months[billing_month]),
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$G:$G)",
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$K:$K)",
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$O:$O)",
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$P:$P)",
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$Q:$Q)",
                f"=SUMIF('Solar Usage'!$A:$A,A{row_number},'Solar Usage'!$R:$R)",
            ]
        )
        summary.cell(row_number, 1).number_format = "mmm yyyy"
        for column in range(3, 8):
            summary.cell(row_number, column).number_format = "#,##0.000"
        summary.cell(row_number, 8).number_format = '₹ #,##0.00'
    _style_sheet(
        summary,
        {"A": 18, "B": 21, "C": 22, "D": 22, "E": 24, "F": 28, "G": 27, "H": 26},
    )
    if months:
        table = Table(
            displayName="SolarMonthlySummary", ref=f"A1:H{summary.max_row}"
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        summary.add_table(table)

    summary.sheet_properties.tabColor = "B18A52"
    usage.sheet_properties.tabColor = "244536"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".solar-consumption-",
            suffix=".xlsx",
            dir=destination.parent,
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
        workbook.save(temp_path)
        os.replace(temp_path, destination)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()
    return destination
